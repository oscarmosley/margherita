"""Exercise actual Swift OOXML writers on Linux. Mac OCR/rendering needs a Mac smoke test."""
import base64
import posixpath
from pathlib import Path
import subprocess
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from test_native_documents_archives import ROOT, SWIFTC, ENV, function

class PDFOfficeLogic(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        source = (ROOT / "Margherita/PDFOffice.swift").read_text()
        pure = 'import Foundation\nenum PDFOffice {\n'
        pure += 'static let relNS = "http://schemas.openxmlformats.org/package/2006/relationships"\n'
        pure += 'static let officeRel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"\n'
        for sig in ["static func supportsFormat(", "static func supports(", "static func xml(", "static func relationships(",
                    "static func types(", "static func docxParts(", "static func xlsxParts(", "static func pptxParts("]:
            pure += function(source, sig) + "\n"
        pure += '}\n'
        pure += r'''
let root = URL(fileURLWithPath: CommandLine.arguments[1])
let texts = ["Hello <&> \" pizza 🍕\n=SUM(A1)\ncontrol\u{0001}character", String(repeating: "猫", count: 70000)]
let sizes = [CGSize(width: 612, height: 792), CGSize(width: 792, height: 612)]
let packages = ["docx": PDFOffice.docxParts(texts, sizes: sizes),
                "xlsx": PDFOffice.xlsxParts(texts), "pptx": PDFOffice.pptxParts(sizes)]
assert(PDFOffice.supports(input: URL(fileURLWithPath: "/tmp/a.pdf"), format: "DOCX"))
assert(!PDFOffice.supports(input: URL(fileURLWithPath: "/tmp/a.txt"), format: "docx"))
for (format, parts) in packages {
    for (path, text) in parts {
        let target = root.appendingPathComponent(format).appendingPathComponent(path)
        try FileManager.default.createDirectory(at: target.deletingLastPathComponent(), withIntermediateDirectories: true)
        try text.write(to: target, atomically: true, encoding: .utf8)
    }
}
'''
        swift = root / "main.swift"
        swift.write_text(pure)
        binary = root / "writer"
        result = subprocess.run([SWIFTC, "-swift-version", "5", str(swift), "-o", str(binary)],
                                env=ENV, capture_output=True, text=True)
        if result.returncode:
            raise RuntimeError(result.stderr)
        cls.packages = root / "packages"
        subprocess.run([str(binary), str(cls.packages)], env=ENV, check=True)
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
        for format, folder in [("docx", "word"), ("pptx", "ppt")]:
            media = cls.packages / format / folder / "media"
            media.mkdir()
            for n in [1, 2]:
                (media / f"page{n}.png").write_bytes(png)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_all_xml_and_relationship_targets(self):
        ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        for directory in self.packages.iterdir():
            for file in directory.rglob("*"):
                if file.is_file() and file.suffix in [".xml", ".rels"]:
                    tree = ET.parse(file)
                    if file.suffix == ".rels":
                        # The root .rels is relative to package root; others to owner directory.
                        base = directory if file.name == ".rels" else file.parent.parent
                        for relation in tree.findall(ns + "Relationship"):
                            self.assertTrue((base / relation.attrib["Target"]).resolve().is_file(),
                                            (file, relation.attrib))
            self.assertTrue((directory / "[Content_Types].xml").is_file())

    def test_docx_text_and_all_page_images(self):
        import docx
        directory = self.packages / "docx"
        package = Path(self.tmp.name) / "document.docx"
        with zipfile.ZipFile(package, "w") as archive:
            for file in directory.rglob("*"):
                if file.is_file():
                    archive.write(file, file.relative_to(directory))
        document = docx.Document(package)
        self.assertEqual(len(document.inline_shapes), 2)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn('Hello <&> " pizza 🍕', text)
        self.assertIn("controlcharacter", text)
        self.assertIn("猫" * 70000, text)
        self.assertIn("Extracted text, page 2", text)

    def test_xlsx_inline_strings_and_long_cell_splitting(self):
        sheet = ET.parse(self.packages / "xlsx/xl/worksheets/sheet1.xml")
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        values = [node.text or "" for node in sheet.findall(".//s:t", ns)]
        self.assertIn("=SUM(A1)", values)
        self.assertEqual(sum(v.count("猫") for v in values), 70000)
        self.assertTrue(all(len(v.encode("utf-16-le")) // 2 <= 32767 for v in values))
        self.assertFalse(sheet.findall(".//s:f", ns))

    def test_pptx_page_count_aspect_and_theme(self):
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        presentation = ET.parse(self.packages / "pptx/ppt/presentation.xml")
        self.assertEqual(len(presentation.findall(".//p:sldId", ns)), 2)
        for n in [1, 2]:
            slide = ET.parse(self.packages / f"pptx/ppt/slides/slide{n}.xml")
            extent = slide.find(".//p:pic/p:spPr/a:xfrm/a:ext", ns)
            self.assertLessEqual(int(extent.attrib["cx"]), 9144000)
            self.assertLessEqual(int(extent.attrib["cy"]), 6858000)
        self.assertTrue((self.packages / "pptx/ppt/theme/theme1.xml").exists())

    def test_ocr_source_safety_contract(self):
        source = (ROOT / "Margherita/NativeOCR.swift").read_text()
        self.assertIn("import Vision", source)
        self.assertIn("setTextDrawingMode(.invisible)", source)
        self.assertIn("kCGImageSourceCreateThumbnailWithTransform: true", source)
        self.assertIn("getDrawingTransform(.mediaBox", source)
        self.assertIn("No readable text", source)
        self.assertIn("CGImageSourceGetCount(source) == 1", source)
        self.assertIn("context?.closePDF()", source)
        self.assertIn("if !searchable, let pdf = pdf", source)
        self.assertIn("let pageText = try text(page: page, language: language)", source)
        self.assertIn("texts.append(pageText)", source)
        self.assertIn("CGImageSourceCopyTypeIdentifiers()", source)
        office = (ROOT / "Margherita/PDFOffice.swift").read_text()
        self.assertIn('if !["docx", "pptx"].contains(format), !texts.contains', office)


@unittest.skipUnless(__import__("sys").platform == "darwin", "Requires macOS Vision/PDFKit SDK")
class NativeOCRMacSmoke(unittest.TestCase):
    def test_real_ocr_pdf_and_office_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "main.swift").write_text(r'''
import Foundation
import AppKit
import CoreText
import PDFKit
import ImageIO
let root = URL(fileURLWithPath: CommandLine.arguments[1])
let png = root.appendingPathComponent("scan.png")
let sourcePDF = root.appendingPathComponent("source.pdf")
let scanPDF = root.appendingPathComponent("scan.pdf")
let ctx = CGContext(data: nil, width: 1400, height: 500, bitsPerComponent: 8, bytesPerRow: 5600,
                    space: CGColorSpaceCreateDeviceRGB(),
                    bitmapInfo: CGImageAlphaInfo.premultipliedLast.rawValue)!
ctx.setFillColor(CGColor(gray: 1, alpha: 1)); ctx.fill(CGRect(x: 0, y: 0, width: 1400, height: 500))
let attrs: [NSAttributedString.Key: Any] = [.font: NSFont.systemFont(ofSize: 72),
                                           .foregroundColor: NSColor.black]
let line = CTLineCreateWithAttributedString(NSAttributedString(string: "MARgherita OCR 12345", attributes: attrs) as CFAttributedString)
ctx.textPosition = CGPoint(x: 60, y: 220); CTLineDraw(line, ctx)
let image = ctx.makeImage()!
let dest = CGImageDestinationCreateWithURL(png as CFURL, "public.png" as CFString, 1, nil)!
CGImageDestinationAddImage(dest, image, nil); assert(CGImageDestinationFinalize(dest))
var box = CGRect(x: 0, y: 0, width: 700, height: 250)
let pdf = CGContext(sourcePDF as CFURL, mediaBox: &box, nil)!
pdf.beginPDFPage(nil)
pdf.textPosition = CGPoint(x: 20, y: 100)
let small = CTLineCreateWithAttributedString(NSAttributedString(string: "Native selectable text 67890", attributes: [.font: NSFont.systemFont(ofSize: 20)]) as CFAttributedString)
CTLineDraw(small, pdf); pdf.endPDFPage(); pdf.closePDF()
let scan = CGContext(scanPDF as CFURL, mediaBox: &box, nil)!
scan.beginPDFPage(nil); scan.draw(image, in: box); scan.endPDFPage(); scan.closePDF()
let mixed = PDFDocument()
mixed.insert(PDFDocument(url: sourcePDF)!.page(at: 0)!, at: 0)
let blank = PDFPage()
blank.setBounds(box, for: .mediaBox)
mixed.insert(blank, at: 1)
mixed.insert(PDFDocument(url: scanPDF)!.page(at: 0)!, at: 2)
let mixedURL = root.appendingPathComponent("mixed.pdf")
assert(mixed.write(to: mixedURL))
try NativeOCR.convert(input: mixedURL, to: root.appendingPathComponent("mixed.txt"), format: "txt")
let mixedText = try String(contentsOf: root.appendingPathComponent("mixed.txt"), encoding: .utf8)
let mixedPages = mixedText.components(separatedBy: "\n\n\u{000C}\n\n")
assert(mixedPages.count == 3)
assert(mixedPages[0].contains("Native selectable text 67890"))
assert(mixedPages[1].trimmingCharacters(in: .whitespacesAndNewlines).isEmpty)
assert(mixedPages[2].contains("12345"))
try NativeOCR.convert(input: png, to: root.appendingPathComponent("ocr.txt"), format: "txt")
let text = try String(contentsOf: root.appendingPathComponent("ocr.txt"), encoding: .utf8)
assert(text.contains("12345"), text)
try NativeOCR.convert(input: scanPDF, to: root.appendingPathComponent("searchable.pdf"), format: "searchable-pdf")
let searchable = PDFDocument(url: root.appendingPathComponent("searchable.pdf"))!
assert(searchable.pageCount == 1)
assert((searchable.string ?? "").contains("12345"))
for format in ["docx", "pptx", "xlsx", "rtf", "txt"] {
    try PDFOffice.convert(input: sourcePDF, to: root.appendingPathComponent("result." + format), format: format)
}
try PDFOffice.convert(input: scanPDF, to: root.appendingPathComponent("scan.txt"), format: "txt")
assert(try String(contentsOf: root.appendingPathComponent("scan.txt"), encoding: .utf8).contains("12345"))
'''.replace('assert(try String(contentsOf: root.appendingPathComponent("scan.txt"), encoding: .utf8).contains("12345"))',
'''let scannedText = try String(contentsOf: root.appendingPathComponent("scan.txt"), encoding: .utf8)
assert(scannedText.contains("12345"))'''))
            binary = root / "smoke"
            sources = ["ConversionCancellation.swift", "NativeArchives.swift", "NativeOCR.swift", "PDFOffice.swift"]
            result = subprocess.run(["/usr/bin/swiftc", "-swift-version", "5",
                                     *[str(ROOT / "Margherita" / f) for f in sources],
                                     str(root / "main.swift"), "-o", str(binary),
                                     "-framework", "AppKit", "-framework", "PDFKit",
                                     "-framework", "Vision", "-framework", "CoreText",
                                     "-framework", "ImageIO"], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            subprocess.run([str(binary), str(root)], check=True, timeout=180)
            for format in ["docx", "pptx", "xlsx"]:
                with zipfile.ZipFile(root / ("result." + format)) as archive:
                    self.assertIsNone(archive.testzip())

if __name__ == "__main__":
    unittest.main()
