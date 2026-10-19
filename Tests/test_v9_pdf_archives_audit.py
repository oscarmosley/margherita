"""Executable portable-route audit. No OCR/rendering claim is made on Linux.
Uses production Swift functions, real zip processes and independent package readers.
"""
import base64
import gzip
import os
from pathlib import Path
import shutil
import subprocess
import tarfile
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from test_native_documents_archives import ROOT, SWIFTC, ENV, function

class V9PortableRouteAudit(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        archive = (ROOT / "Margherita/NativeArchives.swift").read_text()
        if os.uname().sysname != "Darwin":
            if not shutil.which("bsdtar"):
                raise unittest.SkipTest("BSD tar required, GNU tar is not a faithful macOS substitute")
            archive = archive.replace("import Darwin", "import Glibc").replace('"/usr/bin/tar"', '"/usr/bin/bsdtar"')
        (cls.root / "Archives.swift").write_text(archive)
        pdf = (ROOT / "Margherita/PDFOffice.swift").read_text()
        pure = 'import Foundation\nenum PDFOffice {\n'
        pure += 'static let relNS = "http://schemas.openxmlformats.org/package/2006/relationships"\n'
        pure += 'static let officeRel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"\n'
        for sig in ["static func xml(", "static func relationships(", "static func types(",
                    "static func docxParts(", "static func xlsxParts(", "static func pptxParts("]:
            pure += function(pdf, sig) + "\n"
        pure += '}\nenum Documents {\nstatic let fm = FileManager.default\n'
        docs = (ROOT / "Margherita/NativeDocuments.swift").read_text()
        for sig in ["static func xml(", "private static func put(", "private static func writeEPUB("]:
            pure += function(docs, sig).replace("private static", "static") + "\n"
        pure += '}\n'
        (cls.root / "Writers.swift").write_text(pure)
        (cls.root / "main.swift").write_text(r'''
import Foundation
let a = CommandLine.arguments
let root = URL(fileURLWithPath: a[2])
do {
 if a[1] == "extract" {
  try NativeArchives.extract(input: root, to: URL(fileURLWithPath: a[3]))
 } else if a[1] == "archive" {
  try NativeArchives.convert(input: [root], to: URL(fileURLWithPath: a[3]), format: a[4])
 } else {
  let texts = ["First <&> page 🍕\n=1+2", String(repeating: "猫🍕", count: 18000) + " LAST_SENTINEL"]
  let sizes = [CGSize(width: 612, height: 792), CGSize(width: 792, height: 612)]
  let packages = ["fixed.docx": PDFOffice.docxParts(texts, sizes: sizes),
                  "editable.docx": PDFOffice.docxParts(texts, sizes: [], layout: "editable"),
                  "book.xlsx": PDFOffice.xlsxParts(texts), "slides.pptx": PDFOffice.pptxParts(sizes)]
  let png = Data(base64Encoded: "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")!
  for (name, parts) in packages {
   let folder = root.appendingPathComponent(name + "-parts")
   for (path, text) in parts {
    let target = folder.appendingPathComponent(path)
    try FileManager.default.createDirectory(at: target.deletingLastPathComponent(), withIntermediateDirectories: true)
    try text.write(to: target, atomically: true, encoding: .utf8)
   }
   if name == "fixed.docx" || name == "slides.pptx" {
    let media = folder.appendingPathComponent(name == "fixed.docx" ? "word/media" : "ppt/media")
    try FileManager.default.createDirectory(at: media, withIntermediateDirectories: true)
    for id in 1...2 { try png.write(to: media.appendingPathComponent("page\(id).png")) }
   }
   _ = try NativeProcess.run("/usr/bin/zip", ["-q", "-X", "-r", root.appendingPathComponent(name).path, "."], at: folder)
  }
  let work = root.appendingPathComponent("epub-work")
  try FileManager.default.createDirectory(at: work, withIntermediateDirectories: true)
  try Documents.writeEPUB(texts.joined(separator: "\n"), title: "Audit <&>", output: root.appendingPathComponent("book.epub"), work: work)
 }
} catch { fputs(error.localizedDescription + "\n", stderr); exit(1) }
''')
        cls.binary = cls.root / "routes"
        run = subprocess.run([SWIFTC, "-swift-version", "5", str(ROOT / "Margherita/ConversionCancellation.swift"), str(cls.root / "Archives.swift"),
                              str(cls.root / "Writers.swift"), str(cls.root / "main.swift"),
                              "-o", str(cls.binary)], env=ENV, capture_output=True, text=True)
        if run.returncode:
            raise RuntimeError(run.stderr)
        cls.call("packages", cls.root)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    @classmethod
    def call(cls, *args, success=True):
        result = subprocess.run([str(cls.binary), *map(str, args)], env=ENV, capture_output=True, text=True, timeout=120)
        if (result.returncode == 0) != success:
            raise AssertionError(result.stderr)
        return result

    def test_zip_tar_tgz_gz_gzip_roundtrip(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "payload"
            source.mkdir()
            (source / "nested").mkdir()
            data = bytes(range(256)) * 1024
            (source / "nested" / "payload.bin").write_bytes(data)
            (source / "[Content_Types].xml").write_text("<Types/>")
            for ext in ["zip", "tar", "tgz", "gz", "gzip"]:
                archive = root / ("output." + ext)
                self.call("archive", source, archive, ext)
                self.call("extract", archive, root / ext)
                self.assertEqual((root / ext / "payload/nested/payload.bin").read_bytes(), data)
                self.assertEqual((root / ext / "payload/[Content_Types].xml").read_text(), "<Types/>")
                if ext == "zip":
                    with zipfile.ZipFile(archive) as z:
                        self.assertIsNone(z.testzip())
                elif ext == "tar":
                    with tarfile.open(archive) as t:
                        self.assertIn("./payload/nested/payload.bin", t.getnames())
                else:
                    self.assertGreater(len(gzip.decompress(archive.read_bytes())), len(data))

    def test_7z_native_libarchive_fixture(self):
        tar = "/usr/bin/tar" if os.uname().sysname == "Darwin" else shutil.which("bsdtar")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "payload.txt").write_text("Native 7z fixture 猫🍕 END")
            archive = root / "fixture.7z"
            create = subprocess.run([tar, "--format=7zip", "-cf", str(archive),
                                     "-C", str(root), "payload.txt"], capture_output=True)
            if create.returncode:
                self.skipTest("Installed libarchive lacks 7zip writer")
            self.call("extract", archive, root / "extracted")
            self.assertEqual((root / "extracted/payload.txt").read_text(), "Native 7z fixture 猫🍕 END")

    def test_plain_bzip2_xz_streams(self):
        import bz2
        import lzma
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            payload = b"native compression roundtrip " * 8192
            for ext, compress in [("bz2", bz2.compress), ("xz", lzma.compress)]:
                archive = root / ("fixture." + ext)
                archive.write_bytes(compress(payload))
                self.call("extract", archive, root / ext)
                self.assertEqual((root / ext / "fixture").read_bytes(), payload)

    def test_output_archive_rejects_unroundtrippable_names(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "payload"
            source.mkdir()
            (source / "bad*.txt").write_text("must not generate an unreadable archive")
            self.call("archive", source, root / "bad.zip", "zip", success=False)
            self.assertFalse((root / "bad.zip").exists())

    @unittest.skipIf(os.uname().sysname == "Darwin", "Fixture needs case-sensitive filesystem")
    def test_nested_case_collision_rejected_before_creation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "payload"
            source.mkdir()
            (source / "Name.txt").write_text("one")
            (source / "name.txt").write_text("two")
            self.call("archive", source, root / "bad.zip", "zip", success=False)
            self.assertFalse((root / "bad.zip").exists())

    def test_packages_are_readable_and_not_truncated(self):
        import docx
        import openpyxl
        import pptx
        fixed = docx.Document(self.root / "fixed.docx")
        editable = docx.Document(self.root / "editable.docx")
        self.assertEqual(len(fixed.inline_shapes), 2)
        self.assertEqual(len(editable.inline_shapes), 0)
        for document in [fixed, editable]:
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("LAST_SENTINEL", text)
            self.assertEqual(text.count("猫🍕"), 18000)
        workbook = openpyxl.load_workbook(self.root / "book.xlsx")
        values = [str(row[2]) for row in list(workbook.active.values)[1:]]
        self.assertIn("LAST_SENTINEL", "".join(values))
        self.assertEqual("".join(values).count("猫🍕"), 18000)
        self.assertTrue(all(len(v.encode("utf-16-le")) // 2 <= 32767 for v in values))
        presentation = pptx.Presentation(self.root / "slides.pptx")
        self.assertEqual(len(presentation.slides), 2)
        self.assertTrue(all(len(slide.shapes) == 1 for slide in presentation.slides))
        for slide in presentation.slides:
            self.assertGreater(len(slide.shapes[0].image.blob), 0)
        for name in ["fixed.docx", "editable.docx", "book.xlsx", "slides.pptx", "book.epub"]:
            with zipfile.ZipFile(self.root / name) as z:
                self.assertIsNone(z.testzip())
                for member in z.namelist():
                    if member.endswith((".xml", ".rels", ".xhtml", ".opf", ".ncx")):
                        ET.fromstring(z.read(member))
            target = self.root / (name + "-extracted")
            self.call("extract", self.root / name, target)
            self.assertGreater(len(list(target.rglob("*"))), 2)

    def test_epub_mimetype_spine_and_full_text(self):
        with zipfile.ZipFile(self.root / "book.epub") as z:
            self.assertEqual(z.infolist()[0].filename, "mimetype")
            self.assertEqual(z.infolist()[0].compress_type, zipfile.ZIP_STORED)
            self.assertEqual(z.read("mimetype"), b"application/epub+zip")
            opf = ET.fromstring(z.read("OEBPS/content.opf"))
            ns = {"o": "http://www.idpf.org/2007/opf"}
            self.assertEqual(opf.find("o:spine/o:itemref", ns).attrib["idref"], "chapter")
            chapter = ET.fromstring(z.read("OEBPS/chapter.xhtml"))
            text = "".join(chapter.itertext())
            self.assertEqual(text.count("猫🍕"), 18000)
            self.assertIn("LAST_SENTINEL", text)


@unittest.skipUnless(os.uname().sysname == "Darwin", "CoreText/AppKit execution requires Mac")
class V9SlideTextMacAudit(unittest.TestCase):
    def test_actual_coretext_fit_and_continuation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "main.swift").write_text(r'''
import Foundation
import AppKit
import CoreText
import PDFKit
let root = URL(fileURLWithPath: CommandLine.arguments[1])
var media = CGRect(x: 0, y: 0, width: 400, height: 300)
let target = root.appendingPathComponent("overflow.pdf")
let ctx = CGContext(target as CFURL, mediaBox: &media, nil)!
let original = String(repeating: "Recover all these words 12345. ", count: 120) + " END_SENTINEL"
var remaining = original
var pages = 0
while !remaining.isEmpty {
    pages += 1
    assert(pages < 100)
    ctx.beginPDFPage(nil)
    let next = NativeSlides.drawText(remaining, in: CGRect(x: 20, y: 20, width: 160, height: 65), context: ctx)
    assert((next as NSString).length < (remaining as NSString).length)
    remaining = next
    ctx.endPDFPage()
}
ctx.closePDF()
assert(pages > 1)
let pdf = PDFDocument(url: target)!
let recovered = (pdf.string ?? "").filter { !$0.isWhitespace }
assert(recovered == original.filter { !$0.isWhitespace })
assert(recovered.contains("END_SENTINEL"))
''')
            binary = root / "slides"
            subprocess.run(["/usr/bin/swiftc", "-swift-version", "5",
                            str(ROOT / "Margherita/NativeArchives.swift"),
                            str(ROOT / "Margherita/NativeSlides.swift"),
                            str(root / "main.swift"), "-o", str(binary),
                            "-framework", "AppKit", "-framework", "CoreText",
                            "-framework", "PDFKit", "-framework", "ImageIO"],
                           check=True, timeout=120)
            subprocess.run([str(binary), str(root)], check=True, timeout=120)

if __name__ == "__main__":
    unittest.main()
